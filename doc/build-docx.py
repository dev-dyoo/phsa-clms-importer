#!/usr/bin/env python3
"""Build the Word version of the CLMS Data Conversion Report from its Markdown source.

Usage:
    PYTHONPATH=/tmp/pylibs python3 doc/build-docx.py

Reads  doc/CLMS-Data-Conversion-Report-Final.md
Writes doc/CLMS Data Conversion Report - Final.docx

Uses only python-docx default-template built-in styles (Title, Subtitle,
Heading 1-4, Normal, List Bullet, Table Grid, Intense Quote). Supported
Markdown subset:
  <!-- TITLE: ... -->     -> Title paragraph + core property
  <!-- SUBTITLE: ... -->  -> Subtitle paragraph
  <!-- META: K | V -->    -> a 2-column borderless table row (K bold, V plain)
  <!-- PAGEBREAK -->      -> page break
  <!-- TOC -->            -> a Word TOC field (right-click to update)
  # / ## / ### / ####     -> Heading 1-4
  - item                  -> List Bullet
  > quote                 -> Intense Quote
  | a | b |               -> Table Grid (the row before a `---` separator becomes the header)
  inline **bold**, `code`, [text](url), &nbsp;, &lt; &gt; &amp;
"""

import datetime
import os
import re
import sys

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "CLMS-Data-Conversion-Report-Final.md")
OUT = os.path.join(HERE, "CLMS Data Conversion Report - Final.docx")

DOC_TITLE = "CLMS Data Conversion Report - Final"
DOC_AUTHOR = "Attain Solutions Inc."
DOC_DATE = datetime.datetime(2026, 5, 11)

INLINE_RE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")


def deentity(text):
    return (text.replace("&nbsp;", " ")
                .replace("&lt;", "<")
                .replace("&gt;", ">")
                .replace("&amp;", "&"))


def add_inline_runs(paragraph, text):
    """Append `text` to `paragraph`, honouring **bold**, `code`, [text](url)."""
    pos = 0
    for m in INLINE_RE.finditer(text):
        if m.start() > pos:
            paragraph.add_run(deentity(text[pos:m.start()]))
        tok = m.group(0)
        if tok.startswith("**"):
            paragraph.add_run(deentity(tok[2:-2])).bold = True
        elif tok.startswith("`"):
            r = paragraph.add_run(deentity(tok[1:-1]))
            r.font.name = "Consolas"
        else:
            lm = re.match(r"\[([^\]]+)\]\(([^)]+)\)", tok)
            label, url = lm.group(1), lm.group(2)
            paragraph.add_run(deentity(label))
            if url and url != label:
                paragraph.add_run(f" ({deentity(url)})")
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(deentity(text[pos:]))


def style_or(doc, name, fallback="Normal"):
    try:
        doc.styles[name]
        return name
    except KeyError:
        return fallback


def add_page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_toc(doc):
    p = doc.add_paragraph()
    r = p.add_run()._r
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "1-3" \\h \\z \\u '
    sep = OxmlElement("w:fldChar"); sep.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = 'Right-click and choose "Update Field" to build the table of contents.'
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    for el in (begin, instr, sep, placeholder, end):
        r.append(el)


def parse_table_row(line):
    inner = line.strip()
    if inner.startswith("|"):
        inner = inner[1:]
    if inner.endswith("|"):
        inner = inner[:-1]
    return [c.strip() for c in inner.split("|")]


def is_separator_row(cells):
    return any(cells) and all(re.fullmatch(r":?-{2,}:?", (c or "").strip()) for c in cells)


def flush_table(doc, rows):
    if not rows:
        return
    ncols = max(len(c) for c, _ in rows)
    table = doc.add_table(rows=0, cols=ncols)
    table.style = style_or(doc, "Table Grid", None) or table.style
    table.autofit = True
    for cells, is_header in rows:
        cells = cells + [""] * (ncols - len(cells))
        tr = table.add_row().cells
        for i, val in enumerate(cells):
            tr[i].text = ""
            p = tr[i].paragraphs[0]
            add_inline_runs(p, val)
            if is_header:
                for run in p.runs:
                    run.bold = True


def add_meta_table(doc, meta_rows):
    if not meta_rows:
        return
    table = doc.add_table(rows=0, cols=2)
    table.autofit = True
    for k, v in meta_rows:
        tr = table.add_row().cells
        tr[0].text = ""
        kp = tr[0].paragraphs[0]
        add_inline_runs(kp, k)
        for run in kp.runs:
            run.bold = True
        tr[1].text = ""
        add_inline_runs(tr[1].paragraphs[0], v)


def add_heading(doc, style_name, text):
    p = doc.add_paragraph(style=style_name)
    add_inline_runs(p, text)
    return p


def main():
    with open(SRC, "r", encoding="utf-8") as fh:
        lines = fh.read().splitlines()

    doc = Document()
    core = doc.core_properties
    core.title = DOC_TITLE
    core.author = DOC_AUTHOR
    core.last_modified_by = DOC_AUTHOR
    core.created = DOC_DATE
    core.modified = DOC_DATE
    core.revision = 1
    core.category = "Data Conversion Report"
    core.comments = "PHSA - Oracle Fusion Cloud Enterprise Contracts (CLMS) data conversion."

    h_style = {1: style_or(doc, "Heading 1"), 2: style_or(doc, "Heading 2"),
               3: style_or(doc, "Heading 3"), 4: style_or(doc, "Heading 4", "Heading 3")}
    bullet_style = style_or(doc, "List Bullet")
    quote_style = style_or(doc, "Intense Quote", style_or(doc, "Quote"))

    pending_meta = []
    pending_table = []

    def flush_meta():
        nonlocal pending_meta
        if pending_meta:
            add_meta_table(doc, pending_meta)
            pending_meta = []

    def flush_tbl():
        nonlocal pending_table
        if pending_table:
            flush_table(doc, pending_table)
            pending_table = []

    for line in lines:
        stripped = line.strip()

        m = re.match(r"<!--\s*TITLE:\s*(.+?)\s*-->", stripped)
        if m:
            flush_meta(); flush_tbl()
            doc.add_paragraph(deentity(m.group(1)), style=style_or(doc, "Title"))
            continue
        m = re.match(r"<!--\s*SUBTITLE:\s*(.+?)\s*-->", stripped)
        if m:
            flush_meta(); flush_tbl()
            doc.add_paragraph(deentity(m.group(1)), style=style_or(doc, "Subtitle", "Normal"))
            continue
        m = re.match(r"<!--\s*META:\s*(.+?)\s*\|\s*(.+?)\s*-->", stripped)
        if m:
            flush_tbl()
            pending_meta.append((m.group(1), m.group(2)))
            continue
        if re.match(r"<!--\s*PAGEBREAK\s*-->", stripped):
            flush_meta(); flush_tbl()
            add_page_break(doc)
            continue
        if re.match(r"<!--\s*TOC\s*-->", stripped):
            flush_meta(); flush_tbl()
            add_toc(doc)
            continue

        flush_meta()

        if stripped == "":
            flush_tbl()
            continue

        if stripped.startswith("|") and stripped.endswith("|"):
            cells = parse_table_row(stripped)
            if is_separator_row(cells):
                if pending_table:
                    last_cells, _ = pending_table[-1]
                    pending_table[-1] = (last_cells, True)
                continue
            pending_table.append((cells, False))
            continue
        flush_tbl()

        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            level = min(len(m.group(1)), 4)
            add_heading(doc, h_style[level], m.group(2))
            continue

        m = re.match(r"^[-*]\s+(.*)$", stripped)
        if m:
            p = doc.add_paragraph(style=bullet_style)
            add_inline_runs(p, m.group(1))
            continue

        m = re.match(r"^>\s?(.*)$", stripped)
        if m:
            p = doc.add_paragraph(style=quote_style)
            add_inline_runs(p, m.group(1))
            continue

        if re.fullmatch(r"-{3,}|_{3,}|\*{3,}", stripped):
            continue

        p = doc.add_paragraph()
        add_inline_runs(p, stripped)

    flush_meta(); flush_tbl()
    doc.save(OUT)
    print(f"wrote {OUT}")
    print(f"  paragraphs={len(doc.paragraphs)} tables={len(doc.tables)}")


if __name__ == "__main__":
    if not os.path.exists(SRC):
        sys.exit(f"source not found: {SRC}")
    main()
